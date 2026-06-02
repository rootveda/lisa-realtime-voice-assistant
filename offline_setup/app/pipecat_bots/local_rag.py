"""Lexical RAG using SQLite FTS5 — fully offline; index lives under offline_setup/rag_data/."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import time
from pathlib import Path

from loguru import logger

from pipecat_bots.repo_root import resolve_repo_root

_CHUNK = 1200
_CHUNK_OVERLAP = 200

# Images: indexed text = Pillow metadata (dimensions, format); no OCR offline.
RAG_IMAGE_SUFFIXES = frozenset(
    {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".ico"}
)
# Audio/video: indexed placeholder text only (no transcription offline).
RAG_MEDIA_SUFFIXES = frozenset(
    {
        ".mp4",
        ".webm",
        ".mov",
        ".mkv",
        ".avi",
        ".mp3",
        ".wav",
        ".ogg",
        ".m4a",
        ".flac",
        ".opus",
    }
)

_SUPPORTED_SUFFIXES = (
    {".md", ".txt", ".markdown", ".pdf", ".docx", ".csv", ".xlsx", ".xls", ".pptx"}
    | RAG_IMAGE_SUFFIXES
    | RAG_MEDIA_SUFFIXES
)
# Public alias for callers (attachments mirror / UI).
RAG_INDEXABLE_SUFFIXES = _SUPPORTED_SUFFIXES


def _env_positive_int(name: str, default: int) -> int | None:
    """Parse a positive int from env, or None when unlimited (0 / negative / all / none)."""
    if name not in os.environ:
        return default
    raw = os.environ[name].strip().lower()
    if raw in ("", "all", "none", "full"):
        return None
    try:
        n = int(raw)
        return None if n <= 0 else n
    except ValueError:
        return default


def _pdf_max_pages() -> int | None:
    return _env_positive_int("LOCAL_RAG_PDF_MAX_PAGES", 120)


def _pptx_max_slides() -> int | None:
    return _env_positive_int("LOCAL_RAG_PPTX_MAX_SLIDES", 120)


def _xlsx_max_sheets() -> int | None:
    return _env_positive_int("LOCAL_RAG_XLSX_MAX_SHEETS", 24)


def _last_ingest_cache_path() -> Path:
    return _rag_dir() / "last_ingest_by_path.v1.json"


def _ingest_result_status(result: dict) -> str:
    if not result.get("ok"):
        return "failed"
    if result.get("deduped"):
        return "deduped"
    chunks = result.get("chunks")
    if isinstance(chunks, int) and chunks == 0:
        return "empty"
    return "indexed"


def _with_ingest_status(result: dict) -> dict:
    out = dict(result)
    out["status"] = _ingest_result_status(out)
    return out


def _load_last_ingest_results() -> dict[str, dict]:
    path = _last_ingest_cache_path()
    if not path.is_file():
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    entries = raw.get("entries") if isinstance(raw, dict) else None
    if not isinstance(entries, dict):
        return {}
    return {str(k): v for k, v in entries.items() if isinstance(v, dict)}


def _save_last_ingest_results(results: list[dict]) -> None:
    entries: dict[str, dict] = {}
    for r in results:
        pstr = str(r.get("path") or "").strip()
        if not pstr:
            continue
        key = _path_key(pstr)
        entries[key] = {
            "status": r.get("status") or _ingest_result_status(r),
            "error": r.get("error"),
            "chunks": r.get("chunks"),
            "updated_at": int(time.time()),
        }
    path = _last_ingest_cache_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"version": 1, "entries": entries}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _merge_last_ingest_result(result: dict) -> None:
    pstr = str(result.get("path") or "").strip()
    if not pstr:
        return
    cache = _load_last_ingest_results()
    cache[_path_key(pstr)] = {
        "status": result.get("status") or _ingest_result_status(result),
        "error": result.get("error"),
        "chunks": result.get("chunks"),
        "updated_at": int(time.time()),
    }
    path = _last_ingest_cache_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"version": 1, "entries": cache}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _rag_dir() -> Path:
    base = Path(os.environ.get("LOCAL_RAG_DATA_DIR", "")).expanduser()
    if base.is_absolute() and base.parts:
        return base.resolve()
    root = resolve_repo_root()
    return (root / "offline_setup" / "rag_data").resolve()


def rag_db_path() -> Path:
    p = os.environ.get("LOCAL_RAG_DB_PATH", "").strip()
    if p:
        return Path(p).expanduser().resolve()
    return _rag_dir() / "index.sqlite"


def rag_documents_dir() -> Path:
    d = _rag_dir() / "documents"
    d.mkdir(parents=True, exist_ok=True)
    return d


def rag_from_chat_dir() -> Path:
    """Mirrored chat attachments + extracted text for long-term RAG (under rag_data/)."""
    d = _rag_dir() / "from_chat"
    d.mkdir(parents=True, exist_ok=True)
    return d


def rag_dedup_dir() -> Path:
    """Content-addressed blobs (SHA-256) shared by symlinks from mirrors / uploads."""
    d = _rag_dir() / "dedup"
    d.mkdir(parents=True, exist_ok=True)
    return d


def store_dedup_bytes(raw: bytes, suffix: str) -> tuple[Path, bool]:
    """Persist bytes under ``dedup/{aa}/{sha256}{suffix}``. Returns ``(absolute_path, created_new_blob)``."""
    h = hashlib.sha256(raw).hexdigest()
    suf = suffix if suffix.startswith(".") else f".{suffix}" if suffix else ""
    dest = rag_dedup_dir() / h[:2] / f"{h}{suf}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.is_file():
        if dest.stat().st_size != len(raw):
            dest.write_bytes(raw)
            return dest.resolve(), True
        return dest.resolve(), False
    dest.write_bytes(raw)
    return dest.resolve(), True


def symlink_or_copy(link_path: Path, target_abs: Path) -> None:
    """Create ``link_path`` as a symlink to ``target_abs``, or copy if symlinks fail."""
    link_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        if link_path.is_symlink() or link_path.is_file():
            link_path.unlink()
    except OSError:
        pass
    try:
        rel = os.path.relpath(target_abs, start=link_path.parent.resolve())
        link_path.symlink_to(rel)
    except OSError:
        shutil.copy2(target_abs, link_path)


def write_dedup_linked_file(dest_path: Path, raw: bytes, *, suffix: str) -> dict[str, bool | str]:
    """Write ``raw`` into dedup storage and place ``dest_path`` as symlink (or copy)."""
    blob_abs, created = store_dedup_bytes(raw, suffix)
    symlink_or_copy(dest_path.absolute(), blob_abs)
    return {"dedup_blob": str(blob_abs), "new_blob": created}


def _connect() -> sqlite3.Connection:
    rag_db_path().parent.mkdir(parents=True, exist_ok=True)
    # ``timeout`` controls how long sqlite3 waits when the file is locked by another
    # connection. Without it we'd see ``database is locked`` under concurrent writers.
    con = sqlite3.connect(str(rag_db_path()), timeout=30.0)
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA synchronous=NORMAL")
    con.execute("PRAGMA busy_timeout=30000")
    return con


def init_schema() -> None:
    con = _connect()
    try:
        con.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS rag_fts USING fts5(
                path UNINDEXED,
                chunk_idx UNINDEXED,
                body,
                tokenize = 'porter unicode61'
            );
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS rag_doc_content (
                content_sha256 TEXT PRIMARY KEY,
                fts_path TEXT NOT NULL
            );
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS rag_doc_instance (
                path TEXT PRIMARY KEY,
                content_sha256 TEXT NOT NULL
            );
            """
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS rag_doc_instance_sha ON rag_doc_instance(content_sha256)"
        )
        _migrate_legacy_fts_dedup(con)
        con.commit()
    finally:
        con.close()


def _chunk_text(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + _CHUNK, len(text))
        piece = text[start:end]
        chunks.append(piece)
        if end >= len(text):
            break
        start = max(0, end - _CHUNK_OVERLAP)
    return chunks


def _normalize_for_hash(text: str) -> str:
    """Must match the text passed into :func:`_chunk_text` (stable SHA-256 for dedupe)."""
    return text.replace("\r\n", "\n").strip()


def _content_sha256_hex(text: str) -> str:
    return hashlib.sha256(_normalize_for_hash(text).encode("utf-8")).hexdigest()


def _insert_fts_chunks(con: sqlite3.Connection, path_str: str, chunks: list[str]) -> None:
    for i, body in enumerate(chunks):
        con.execute(
            "INSERT INTO rag_fts(path, chunk_idx, body) VALUES (?, ?, ?)",
            (path_str, i, body),
        )


def _promote_canonical_for_sha(con: sqlite3.Connection, content_sha256: str) -> bool:
    """Rebuild FTS for another path that shares ``content_sha256`` when the canonical file goes away."""
    rows = con.execute(
        "SELECT path FROM rag_doc_instance WHERE content_sha256 = ? ORDER BY path",
        (content_sha256,),
    ).fetchall()
    for (pstr,) in rows:
        p = Path(pstr)
        if not p.is_file():
            continue
        text, err = _extract_text_from_path(p)
        if err:
            continue
        text = text or ""
        chunks = _chunk_text(text)
        if not chunks:
            continue
        con.execute("DELETE FROM rag_fts WHERE path = ?", (pstr,))
        _insert_fts_chunks(con, pstr, chunks)
        con.execute(
            "INSERT OR REPLACE INTO rag_doc_content(content_sha256, fts_path) VALUES (?, ?)",
            (content_sha256, pstr),
        )
        logger.info(
            "[local_rag] promoted canonical FTS path=%s sha=%s…",
            pstr,
            content_sha256[:12],
        )
        return True
    return False


def _detach_path_and_maybe_promote(con: sqlite3.Connection, pstr: str) -> None:
    """Remove FTS rows and registry entry for ``pstr``; promote another path if this was canonical."""
    con.execute("DELETE FROM rag_fts WHERE path = ?", (pstr,))
    row = con.execute(
        "SELECT content_sha256 FROM rag_doc_instance WHERE path = ?",
        (pstr,),
    ).fetchone()
    con.execute("DELETE FROM rag_doc_instance WHERE path = ?", (pstr,))
    if not row:
        return
    sha = row[0]
    crow = con.execute(
        "SELECT fts_path FROM rag_doc_content WHERE content_sha256 = ?",
        (sha,),
    ).fetchone()
    if crow and crow[0] == pstr:
        con.execute("DELETE FROM rag_doc_content WHERE content_sha256 = ?", (sha,))
        if not _promote_canonical_for_sha(con, sha):
            logger.debug(
                "[local_rag] no promote candidate for sha=%s… after dropping %s",
                sha[:12],
                pstr,
            )


def _migrate_legacy_fts_dedup(con: sqlite3.Connection) -> None:
    """One-time: collapse duplicate FTS rows when ``rag_doc_instance`` is empty but legacy FTS has rows."""
    n_inst = int(con.execute("SELECT COUNT(*) FROM rag_doc_instance").fetchone()[0])
    if n_inst > 0:
        return
    n_fts = int(con.execute("SELECT COUNT(*) FROM rag_fts").fetchone()[0])
    if n_fts == 0:
        return
    paths = [r[0] for r in con.execute("SELECT DISTINCT path FROM rag_fts").fetchall()]
    by_sha: dict[str, list[str]] = {}
    for pstr in paths:
        p = Path(pstr)
        if not p.is_file():
            continue
        text, err = _extract_text_from_path(p)
        if err or not text:
            continue
        sha = _content_sha256_hex(text)
        by_sha.setdefault(sha, []).append(pstr)
    merged = 0
    for sha, plist in by_sha.items():
        plist = sorted(plist)
        canonical = plist[0]
        for dup in plist[1:]:
            con.execute("DELETE FROM rag_fts WHERE path = ?", (dup,))
            con.execute(
                "INSERT OR REPLACE INTO rag_doc_instance(path, content_sha256) VALUES (?, ?)",
                (dup, sha),
            )
            merged += 1
        con.execute(
            "INSERT OR REPLACE INTO rag_doc_content(content_sha256, fts_path) VALUES (?, ?)",
            (sha, canonical),
        )
        con.execute(
            "INSERT OR REPLACE INTO rag_doc_instance(path, content_sha256) VALUES (?, ?)",
            (canonical, sha),
        )
    if merged:
        logger.info(
            "[local_rag] legacy migration: merged %s duplicate FTS paths by content hash",
            merged,
        )


def _fts_match_query(user_q: str) -> str:
    """Build conservative FTS5 AND query from user text."""
    raw = re.sub(r'[^\w\s\-]', ' ', user_q, flags=re.UNICODE)
    parts = [p for p in raw.split() if p][:16]
    if not parts:
        return ""
    return " AND ".join(f"{p}" for p in parts)


def describe_image_from_bytes(raw: bytes, fname: str) -> str:
    """Return searchable text for an image: dimensions, format, mode (Pillow). No OCR."""
    try:
        from io import BytesIO

        from PIL import Image

        im = Image.open(BytesIO(raw))
        im.load()
        w, h = im.size
        fmt = (im.format or "?").upper()
        mode = im.mode or "?"
        lines = [
            f"[Image: {Path(fname).name}]",
            f"Dimensions: {w} x {h} pixels; format={fmt}; mode={mode}.",
            "Visual content is not auto-described in the offline RAG/attachment path (use Assistant vision or paste a description).",
        ]
        try:
            exif = im.getexif()
            if exif and len(exif) > 0:
                lines.append(f"EXIF: {len(exif)} tag(s) present (not expanded).")
        except Exception:
            pass
        return "\n".join(lines)
    except Exception as e:
        return (
            f"[Image: {Path(fname).name}]\n"
            f"Could not read image metadata ({e}). Size: {len(raw)} bytes."
        )


def describe_media_binary_from_bytes(fname: str, raw: bytes, mime: str) -> str:
    """Placeholder text for audio/video and other media without speech-to-text offline."""
    mt = (mime or "").strip() or "application/octet-stream"
    return (
        f"[Media: {Path(fname).name}]\n"
        f"MIME: {mt}; size: {len(raw)} bytes.\n"
        "This stack does not transcribe audio or video for RAG. Add a separate text note or use an external transcript."
    )


def _extract_text_from_path(abs_path: Path) -> tuple[str | None, str | None]:
    """Return (text, error). Supports docs, Pillow metadata for images, placeholder text for media."""
    suf = abs_path.suffix.lower()
    try:
        if suf in {".md", ".txt", ".markdown"}:
            return abs_path.read_text(encoding="utf-8", errors="replace"), None

        if suf == ".pdf":
            try:
                from pypdf import PdfReader
            except Exception:
                return None, "pypdf not installed for PDF extraction"
            r = PdfReader(str(abs_path))
            pages = r.pages
            cap = _pdf_max_pages()
            if cap is not None:
                pages = pages[:cap]
            parts = []
            for page in pages:
                t = page.extract_text() or ""
                if t.strip():
                    parts.append(t)
            return ("\n\n".join(parts) if parts else ""), None

        if suf == ".docx":
            try:
                from docx import Document
            except Exception:
                return None, "python-docx not installed for DOCX extraction"
            d = Document(str(abs_path))
            out = [p.text for p in d.paragraphs if (p.text or "").strip()]
            return ("\n".join(out) if out else ""), None

        if suf == ".csv":
            lines = []
            with abs_path.open("r", encoding="utf-8", errors="replace", newline="") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(" | ".join(c.strip() for c in row))
            return ("\n".join(lines) if lines else ""), None

        if suf == ".xlsx":
            try:
                from openpyxl import load_workbook
            except Exception:
                return None, "openpyxl not installed for XLSX extraction"
            wb = load_workbook(filename=str(abs_path), read_only=True, data_only=True)
            sheets = wb.worksheets
            cap = _xlsx_max_sheets()
            if cap is not None:
                sheets = sheets[:cap]
            lines = []
            for ws in sheets:
                lines.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [("" if v is None else str(v)).strip() for v in row]
                    if any(vals):
                        lines.append(" | ".join(vals))
            return ("\n".join(lines) if lines else ""), None

        if suf == ".xls":
            try:
                import xlrd
            except Exception:
                return None, "xlrd not installed for XLS extraction"
            wb = xlrd.open_workbook(str(abs_path), on_demand=True)
            names = wb.sheet_names()
            cap = _xlsx_max_sheets()
            if cap is not None:
                names = names[:cap]
            lines = []
            for sname in names:
                ws = wb.sheet_by_name(sname)
                lines.append(f"# Sheet: {sname}")
                for rix in range(ws.nrows):
                    vals = [str(ws.cell_value(rix, cix)).strip() for cix in range(ws.ncols)]
                    if any(v for v in vals):
                        lines.append(" | ".join(vals))
            return ("\n".join(lines) if lines else ""), None

        if suf == ".pptx":
            try:
                from pptx import Presentation
            except Exception:
                return None, "python-pptx not installed for PPTX extraction"
            prs = Presentation(str(abs_path))
            slides = list(prs.slides)
            cap = _pptx_max_slides()
            if cap is not None:
                slides = slides[:cap]
            lines = []
            for i, slide in enumerate(slides, start=1):
                lines.append(f"# Slide {i}")
                for shape in slide.shapes:
                    txt = getattr(shape, "text", "") or ""
                    if txt.strip():
                        lines.append(txt.strip())
            return ("\n".join(lines) if lines else ""), None

        if suf in RAG_IMAGE_SUFFIXES:
            raw = abs_path.read_bytes()
            return describe_image_from_bytes(raw, abs_path.name), None

        if suf in RAG_MEDIA_SUFFIXES:
            raw = abs_path.read_bytes()
            return describe_media_binary_from_bytes(abs_path.name, raw, ""), None
    except OSError as e:
        return None, str(e)
    except Exception as e:
        return None, str(e)
    return None, f"unsupported file type: {suf}"


def ingest_file(abs_path: Path) -> dict:
    """Load one supported document into FTS (replace rows for that path).

    Duplicate content (same extracted text after normalization) shares one FTS row set:
    only the canonical path holds ``rag_fts`` rows; other paths are recorded in
    ``rag_doc_instance`` and reported as indexed without duplicating chunks.
    """
    init_schema()
    abs_path = abs_path.expanduser().absolute()
    pstr = str(abs_path)
    if abs_path.suffix.lower() not in _SUPPORTED_SUFFIXES:
        out = _with_ingest_status(
            {"ok": False, "error": f"unsupported file type: {abs_path.suffix.lower()}", "path": pstr}
        )
        _merge_last_ingest_result(out)
        return out
    text, err = _extract_text_from_path(abs_path)
    if err:
        out = _with_ingest_status({"ok": False, "error": err, "path": pstr})
        _merge_last_ingest_result(out)
        return out
    text = text or ""
    chunks = _chunk_text(text)
    if not chunks:
        con = _connect()
        try:
            _detach_path_and_maybe_promote(con, pstr)
            con.commit()
        finally:
            con.close()
        out = _with_ingest_status({"ok": True, "path": pstr, "chunks": 0})
        _merge_last_ingest_result(out)
        return out

    sha = _content_sha256_hex(text)
    con = _connect()
    try:
        _detach_path_and_maybe_promote(con, pstr)

        row = con.execute(
            "SELECT fts_path FROM rag_doc_content WHERE content_sha256 = ?",
            (sha,),
        ).fetchone()
        if row:
            cand = row[0]
            if cand != pstr and Path(cand).is_file():
                con.execute(
                    "INSERT OR REPLACE INTO rag_doc_instance(path, content_sha256) VALUES (?, ?)",
                    (pstr, sha),
                )
                con.commit()
                out = _with_ingest_status(
                    {
                        "ok": True,
                        "path": pstr,
                        "chunks": len(chunks),
                        "deduped": True,
                        "canonical_path": cand,
                        "content_sha256": sha,
                    }
                )
                _merge_last_ingest_result(out)
                return out
            if cand != pstr and not Path(cand).is_file():
                con.execute("DELETE FROM rag_doc_content WHERE content_sha256 = ?", (sha,))

        _insert_fts_chunks(con, pstr, chunks)
        con.execute(
            "INSERT OR REPLACE INTO rag_doc_content(content_sha256, fts_path) VALUES (?, ?)",
            (sha, pstr),
        )
        con.execute(
            "INSERT OR REPLACE INTO rag_doc_instance(path, content_sha256) VALUES (?, ?)",
            (pstr, sha),
        )
        con.commit()
    finally:
        con.close()
    out = _with_ingest_status(
        {
            "ok": True,
            "path": pstr,
            "chunks": len(chunks),
            "deduped": False,
            "content_sha256": sha,
        }
    )
    _merge_last_ingest_result(out)
    return out


def ingest_documents_dir() -> dict:
    """Ingest all supported docs under ``documents/`` and ``from_chat/``."""
    results = []
    nchunk = 0
    counts = {"indexed": 0, "deduped": 0, "empty": 0, "failed": 0}
    issues: list[dict] = []
    for base in (rag_documents_dir(), rag_from_chat_dir()):
        for f in sorted(base.rglob("*")):
            if not f.is_file():
                continue
            if f.suffix.lower() not in _SUPPORTED_SUFFIXES:
                continue
            r = ingest_file(f)
            results.append(r)
            st = r.get("status") or _ingest_result_status(r)
            counts[st] = counts.get(st, 0) + 1
            if st in ("failed", "empty"):
                issues.append(
                    {
                        "path": r.get("path"),
                        "status": st,
                        "error": r.get("error") if st == "failed" else "no extractable text",
                    }
                )
            if r.get("ok") and isinstance(r.get("chunks"), int):
                nchunk += r["chunks"]
    _save_last_ingest_results(results)
    return {
        "ok": True,
        "files": len(results),
        "total_chunks": nchunk,
        "indexed": counts.get("indexed", 0),
        "deduped": counts.get("deduped", 0),
        "empty": counts.get("empty", 0),
        "failed": counts.get("failed", 0),
        "issues": issues[:40],
        "details": results[:80],
    }


def _is_relative_to(path: Path, base: Path) -> bool:
    try:
        path.absolute().relative_to(base.absolute())
        return True
    except Exception:
        return False


def _catalog_path_str(p: Path | str) -> str:
    """Absolute path under documents/from_chat without following symlinks."""
    return str(Path(p).expanduser().absolute())


def _catalog_paths_for_remove(p: Path) -> list[Path]:
    """Map a UI/delete path to catalog file(s) under documents/ or from_chat/."""
    p = p.expanduser().absolute()
    docs = rag_documents_dir().absolute()
    chat = rag_from_chat_dir().absolute()
    dedup = rag_dedup_dir().absolute()

    if _is_relative_to(p, docs) or _is_relative_to(p, chat):
        return [p]

    if not _is_relative_to(p, dedup):
        return []

    try:
        target = p.resolve()
    except OSError:
        return []

    found: list[Path] = []
    seen: set[str] = set()
    for base in (docs, chat):
        for f in base.rglob("*"):
            if f.is_dir():
                continue
            try:
                if f.is_symlink():
                    if not f.exists() or f.resolve() != target:
                        continue
                elif f.absolute() != p:
                    continue
            except OSError:
                continue
            key = _catalog_path_str(f)
            if key not in seen:
                seen.add(key)
                found.append(f.absolute())
    return found


def _remove_one_catalog_file(p: Path) -> dict:
    """Remove one catalog file/symlink and detach it from index metadata."""
    if p.is_dir():
        return {"ok": False, "error": "path points to a directory; file expected", "path": _catalog_path_str(p)}

    pstr = _catalog_path_str(p)
    init_schema()
    con = _connect()
    try:
        _detach_path_and_maybe_promote(con, pstr)
        # Legacy rows may reference the resolved dedup target instead of the catalog path.
        try:
            if p.is_symlink():
                _detach_path_and_maybe_promote(con, str(p.resolve()))
        except OSError:
            pass
        con.commit()
    finally:
        con.close()

    removed_disk = False
    if p.is_symlink() or p.is_file():
        try:
            p.unlink()
            removed_disk = True
        except OSError as e:
            return {"ok": False, "error": f"failed to remove file: {e}", "path": pstr}
    return {"ok": True, "path": pstr, "removed_disk": removed_disk}


def remove_rag_file(path: str) -> dict:
    """Remove one RAG file from disk (documents/from_chat) and detach it from index metadata."""
    raw = Path(str(path or "")).expanduser().absolute()
    targets = _catalog_paths_for_remove(raw)
    if not targets:
        return {"ok": False, "error": "path is outside allowed RAG roots", "path": str(raw)}

    if len(targets) == 1:
        return _remove_one_catalog_file(targets[0])

    removed: list[str] = []
    errors: list[str] = []
    for t in targets:
        out = _remove_one_catalog_file(t)
        if out.get("ok"):
            removed.append(str(out.get("path") or _catalog_path_str(t)))
        else:
            errors.append(str(out.get("error") or "remove failed"))
    if errors and not removed:
        return {"ok": False, "error": "; ".join(errors), "path": str(raw)}
    return {
        "ok": len(errors) == 0,
        "path": str(raw),
        "removed_catalog_paths": removed,
        "errors": errors or None,
    }


def search_knowledge_base(query: str, *, limit: int = 8) -> dict:
    """Return snippets for LLM tool results."""
    init_schema()
    mq = _fts_match_query(query)
    if not mq:
        return {"ok": True, "hits": [], "note": "empty query"}
    con = _connect()
    try:
        cur = con.execute(
            """
            SELECT path, chunk_idx,
                   snippet(rag_fts, 2, '[', ']', ' … ', 48) AS snip
            FROM rag_fts
            WHERE rag_fts MATCH ?
            ORDER BY rank
            LIMIT ?
            """,
            (mq, max(1, min(limit, 24))),
        )
        rows = cur.fetchall()
    except sqlite3.OperationalError as e:
        logger.warning(f"[local_rag] search error: {e}")
        return {"ok": False, "error": str(e), "hits": []}
    finally:
        con.close()
    hits = [{"path": r[0], "chunk": r[1], "snippet": r[2]} for r in rows]
    return {"ok": True, "hits": hits, "query": query}


def search_indexed_files(query: str, *, limit: int = 12) -> dict:
    """Search file names/paths in the indexed corpus (metadata lookup)."""
    q = (query or "").strip().lower()
    if not q:
        return {"ok": True, "hits": [], "note": "empty query"}
    # Support natural-language requests like "check file named TOY".
    tokens = [t for t in re.findall(r"[a-z0-9_\-\.]+", q) if len(t) >= 2]
    stop = {
        "can",
        "you",
        "check",
        "file",
        "files",
        "name",
        "named",
        "called",
        "the",
        "a",
        "an",
        "for",
        "with",
        "find",
        "search",
        "indexed",
    }
    tokens = [t for t in tokens if t not in stop]
    if not tokens:
        tokens = [q]
    idx = _indexed_paths_set()
    scored = []
    for p in _iter_indexable_files():
        ps = str(p)
        if ps not in idx:
            continue
        name_l = p.name.lower()
        path_l = ps.lower()
        exact = q in name_l or q in path_l
        token_hits = sum(1 for t in tokens if (t in name_l or t in path_l))
        if exact or token_hits > 0:
            score = (100 if exact else 0) + token_hits
            scored.append((score, {"path": ps, "name": p.name, "ext": p.suffix.lower() or "(none)"}))
    scored.sort(key=lambda x: (-x[0], x[1]["name"]))
    lim = max(1, min(limit, 100))
    rows = [row for _, row in scored[:lim]]
    return {"ok": True, "hits": rows, "query": query, "tokens": tokens}


def read_indexed_file_excerpt(query: str, *, max_chars: int = 4000) -> dict:
    """Read a safe excerpt from an indexed file by filename/path fragment."""
    q = (query or "").strip().lower()
    if not q:
        return {"ok": False, "error": "empty query"}
    idx = _indexed_paths_set()
    target: Path | None = None
    for p in _iter_indexable_files():
        ps = str(p)
        if ps not in idx:
            continue
        if q in p.name.lower() or q in ps.lower():
            target = p
            break
    if target is None:
        return {"ok": False, "error": "no indexed file matched query", "query": query}
    text, err = _extract_text_from_path(target)
    if err:
        return {"ok": False, "error": err, "path": str(target)}
    body = (text or "").strip()
    if not body:
        return {"ok": True, "path": str(target), "excerpt": "", "note": "file has no extractable text"}
    lim = max(200, min(int(max_chars), 20000))
    return {"ok": True, "path": str(target), "excerpt": body[:lim]}


def _indexed_paths_set() -> set[str]:
    init_schema()
    con = _connect()
    try:
        rows = con.execute("SELECT path FROM rag_doc_instance").fetchall()
        out: set[str] = set()
        for r in rows:
            if not r or not r[0]:
                continue
            out.add(_catalog_path_str(r[0]))
            try:
                out.add(_path_key(r[0]))
            except OSError:
                pass
        if not out:
            rows = con.execute("SELECT DISTINCT path FROM rag_fts").fetchall()
            for r in rows:
                if not r or not r[0]:
                    continue
                out.add(_catalog_path_str(r[0]))
                try:
                    out.add(_path_key(r[0]))
                except OSError:
                    pass
        return out
    finally:
        con.close()


def _iter_indexable_files() -> list[Path]:
    """All supported files on disk under documents/ and from_chat/ (catalog paths, not dedup targets)."""
    out: list[Path] = []
    for base in (rag_documents_dir(), rag_from_chat_dir()):
        for f in base.rglob("*"):
            if f.is_dir():
                continue
            if f.suffix.lower() not in _SUPPORTED_SUFFIXES:
                continue
            if f.is_symlink() and not f.exists():
                continue
            if not (f.is_file() or f.is_symlink()):
                continue
            out.append(f.expanduser().absolute())
    seen: set[str] = set()
    uniq: list[Path] = []
    for p in sorted(out):
        k = _catalog_path_str(p)
        if k not in seen:
            seen.add(k)
            uniq.append(p)
    return uniq


def _path_key(p: Path | str) -> str:
    return str(Path(p).expanduser().resolve())


def _diagnose_unindexed_file(path: Path) -> tuple[str, str | None]:
    """Probe disk to explain why a file is not in the FTS index."""
    text, err = _extract_text_from_path(path)
    if err:
        return "failed", err
    if not _chunk_text(text or ""):
        return "empty", "no extractable text"
    return "pending", None


def _file_row_status(path: Path, indexed: set[str], last_ingest: dict[str, dict]) -> tuple[str, str | None]:
    pstr = _catalog_path_str(path)
    resolved = _path_key(path)
    if pstr in indexed or resolved in indexed:
        return "chunked", None
    rec = last_ingest.get(pstr)
    if not rec:
        for alt in (str(path), str(path.absolute())):
            rec = last_ingest.get(alt)
            if rec:
                break
    if rec:
        st = str(rec.get("status") or "")
        if st == "failed":
            return "failed", rec.get("error") or "indexing failed"
        if st == "empty":
            return "empty", rec.get("error") or "no extractable text"
    return _diagnose_unindexed_file(path)


def rag_files_detail(*, limit: int = 1000) -> dict:
    """Detailed per-file listing for RAG Arena UI."""
    lim = max(1, min(int(limit), 5000))
    indexed = _indexed_paths_set()
    last_ingest = _load_last_ingest_results()
    rows = []
    for p in _iter_indexable_files()[:lim]:
        pstr = _catalog_path_str(p)
        try:
            st = p.stat()
            sz = int(st.st_size)
            mt = int(st.st_mtime)
        except OSError:
            sz = 0
            mt = 0
        status, reason = _file_row_status(p, indexed, last_ingest)
        rows.append(
            {
                "path": pstr,
                "name": p.name,
                "ext": p.suffix.lower() or "(none)",
                "bytes": sz,
                "mtime": mt,
                "indexed": status == "chunked",
                "status": status,
                "reason": reason,
            }
        )
    rows.sort(key=lambda r: (r["indexed"], r["mtime"]), reverse=True)
    return {"ok": True, "count": len(rows), "limit": lim, "files": rows}


def rag_detail_status() -> dict:
    """Filesystem + FTS stats for UI (Assistant Console / RAG Arena)."""
    init_schema()
    con = _connect()
    try:
        chunks_total = con.execute("SELECT count(*) FROM rag_fts").fetchone()[0]
    finally:
        con.close()

    indexed = _indexed_paths_set()
    disk_files = _iter_indexable_files()
    total_bytes = 0
    ext_counts: dict[str, int] = {}
    files_indexed = 0
    files_failed = 0
    files_empty = 0
    files_pending = 0
    issues_preview: list[dict] = []
    last_ingest = _load_last_ingest_results()
    for p in disk_files:
        try:
            total_bytes += p.stat().st_size
        except OSError:
            pass
        suf = p.suffix.lower() or "(none)"
        ext_counts[suf] = ext_counts.get(suf, 0) + 1
        pstr = _catalog_path_str(p)
        if pstr in indexed:
            files_indexed += 1
            continue
        status, reason = _file_row_status(p, indexed, last_ingest)
        if status == "failed":
            files_failed += 1
            if len(issues_preview) < 20:
                issues_preview.append({"path": pstr, "status": status, "error": reason})
        elif status == "empty":
            files_empty += 1
            if len(issues_preview) < 20:
                issues_preview.append({"path": pstr, "status": status, "error": reason})
        else:
            files_pending += 1

    return {
        "db_path": str(rag_db_path()),
        "documents_dir": str(rag_documents_dir()),
        "chunks_total": chunks_total,
        "from_chat_dir": str(rag_from_chat_dir()),
        "files_on_disk": len(disk_files),
        "files_chunked": files_indexed,
        "files_failed": files_failed,
        "files_empty": files_empty,
        "files_pending": max(0, files_pending),
        "total_bytes": total_bytes,
        "extensions": ext_counts,
        "rag_dir": str(_rag_dir()),
        "issues_preview": issues_preview,
        "files_detail_preview": rag_files_detail(limit=50).get("files", []),
    }


def rag_status() -> dict:
    """Backward-compatible short status + detail fields used by newer UIs."""
    d = rag_detail_status()
    return {
        "db_path": d["db_path"],
        "documents_dir": d["documents_dir"],
        "chunks": d["chunks_total"],
        "from_chat_dir": d.get("from_chat_dir"),
        "chunks_total": d["chunks_total"],
        "files_on_disk": d["files_on_disk"],
        "files_chunked": d["files_chunked"],
        "files_failed": d.get("files_failed", 0),
        "files_empty": d.get("files_empty", 0),
        "files_pending": d["files_pending"],
        "issues_preview": d.get("issues_preview", []),
        "total_bytes": d["total_bytes"],
        "extensions": d["extensions"],
        "rag_dir": d["rag_dir"],
    }


def augment_voice_transcript_with_rag(user_text: str) -> str:
    """Prepend offline RAG snippets to STT text for voice / voice+vision (env-gated)."""
    raw = (user_text or "").strip()
    if not raw:
        return user_text or ""
    if os.environ.get("VOICE_LOCAL_RAG", "1").strip().lower() in ("0", "false", "no"):
        return user_text
    try:
        lim = int(os.environ.get("VOICE_RAG_HIT_LIMIT", "6"))
    except ValueError:
        lim = 6
    lim = max(1, min(lim, 12))
    r = search_knowledge_base(raw, limit=lim)
    lines = []
    if r.get("ok") and r.get("hits"):
        for h in r["hits"][:lim]:
            p = h.get("path", "")
            sn = (h.get("snippet") or "").replace("\n", " ")
            lines.append(f"- ({p}) {sn}")
    else:
        # Voice path often asks by filename (e.g., "check file TOY").
        # Fall back to indexed file-name metadata search so assistant can answer reliably.
        f = search_indexed_files(raw, limit=lim)
        for hit in (f.get("hits") or [])[:lim]:
            p = hit.get("path", "")
            nm = hit.get("name", "")
            lines.append(f"- indexed file match: {nm} ({p})")
    if not lines:
        return user_text
    block = "\n".join(lines)
    max_chars = int(os.environ.get("VOICE_RAG_CONTEXT_CHARS", "6000"))
    if len(block) > max_chars:
        block = block[: max_chars - 3] + "..."
    return f"[Offline knowledge base — snippets]\n{block}\n\nUser said: {user_text}"
